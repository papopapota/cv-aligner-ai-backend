import io

import docx
import pytest
from pypdf import PdfWriter

from src.domain.errors import InvalidUploadError, UnsupportedFormatError
from src.infrastructure.adapters.out.document_parser_adapter import (
    DocumentParserAdapter,
)


def _normalize(text: str) -> str:
    return " ".join(text.split())


def _blank_pdf_bytes() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    stream = io.BytesIO()
    writer.write(stream)
    return stream.getvalue()


def _docx_bytes(*paragraphs: str, table_cells: tuple[str, ...] = ()) -> bytes:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_cells:
        table = document.add_table(rows=1, cols=len(table_cells))
        for cell, value in zip(table.rows[0].cells, table_cells):
            cell.text = value
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


@pytest.fixture
def parser() -> DocumentParserAdapter:
    return DocumentParserAdapter()


async def test_parses_the_paragraphs_of_a_docx(parser: DocumentParserAdapter) -> None:
    content = _docx_bytes("Ana Developer", "5 years of experience")

    text = await parser.parse("cv.docx", content)

    assert text == "Ana Developer\n5 years of experience"


async def test_includes_the_table_cells_of_a_docx(parser: DocumentParserAdapter) -> None:
    content = _docx_bytes("Ana Developer", table_cells=("Python", "FastAPI"))

    text = await parser.parse("cv.docx", content)

    assert text == "Ana Developer\nPython\nFastAPI"


async def test_skips_blank_paragraphs(parser: DocumentParserAdapter) -> None:
    content = _docx_bytes("Ana Developer", "   ")

    text = await parser.parse("cv.docx", content)

    assert text == "Ana Developer"


async def test_parses_a_docx_regardless_of_the_extension_case(
    parser: DocumentParserAdapter,
) -> None:
    content = _docx_bytes("Ana Developer")

    text = await parser.parse("CV.DOCX", content)

    assert text == "Ana Developer"


async def test_parses_the_sample_pdf(
    parser: DocumentParserAdapter,
    sample_cv_pdf: bytes,
    sample_cv_expected_text: str,
) -> None:
    text = await parser.parse("cv.pdf", sample_cv_pdf)

    assert _normalize(text) == _normalize(sample_cv_expected_text)


async def test_parses_a_pdf_without_a_text_layer_as_empty(
    parser: DocumentParserAdapter,
) -> None:
    text = await parser.parse("cv.pdf", _blank_pdf_bytes())

    assert text == ""


async def test_rejects_a_corrupted_pdf(parser: DocumentParserAdapter) -> None:
    with pytest.raises(InvalidUploadError):
        await parser.parse("cv.pdf", b"this is not a pdf at all")


async def test_rejects_a_corrupted_docx(parser: DocumentParserAdapter) -> None:
    with pytest.raises(InvalidUploadError):
        await parser.parse("cv.docx", b"not a zip archive")


async def test_rejects_an_unsupported_extension(parser: DocumentParserAdapter) -> None:
    with pytest.raises(UnsupportedFormatError):
        await parser.parse("cv.txt", b"plain text")