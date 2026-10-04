import io

import docx
import pytest
from httpx import ASGITransport, AsyncClient
from pypdf import PdfWriter

from src.application.ports.out.cv_parser_port import CVParserPort
from src.application.services.upload_cv_service import UploadCVService
from src.domain.upload_rules import MAX_UPLOAD_SIZE_BYTES
from src.infrastructure.adapters.inbound.api.cv_controller import (
    get_upload_cv_use_case,
)
from src.main import app

_PARSED_TEXT = "Ana Developer\nPython developer with FastAPI experience."
_JOB_DESCRIPTION_TEXT = "We are looking for a Python developer with FastAPI experience."
_DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_TXT_CONTENT_TYPE = "text/plain"


class _StubParser(CVParserPort):
    async def parse(self, filename: str, content: bytes) -> str:
        if filename.endswith(".txt"):
            return _JOB_DESCRIPTION_TEXT
        return _PARSED_TEXT


def _docx_bytes() -> bytes:
    document = docx.Document()
    document.add_paragraph("Ana Developer")
    document.add_paragraph("Python developer with FastAPI experience.")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _txt_bytes(text: str = _JOB_DESCRIPTION_TEXT) -> bytes:
    return text.encode("utf-8")


def _build_client() -> AsyncClient:
    return AsyncClient(transport=ASGITransport(app=app), base_url="http://test")


def _blank_pdf_bytes() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    stream = io.BytesIO()
    writer.write(stream)
    return stream.getvalue()


@pytest.fixture(autouse=True)
def _stub_the_use_case():
    app.dependency_overrides[get_upload_cv_use_case] = lambda: UploadCVService(
        _StubParser()
    )
    yield
    app.dependency_overrides.clear()


@pytest.fixture
def client() -> AsyncClient:
    return _build_client()


async def test_upload_returns_the_parsed_cv_and_job_description(client: AsyncClient) -> None:
    response = await client.post(
        "/cv/upload",
        files={
            "cv_file": ("cv.docx", _docx_bytes(), _DOCX_CONTENT_TYPE),
            "job_description_file": ("job.txt", _txt_bytes(), _TXT_CONTENT_TYPE),
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["filename"] == "cv.docx"
    assert payload["content_type"] == _DOCX_CONTENT_TYPE
    assert payload["extracted_text"] == _PARSED_TEXT
    assert payload["characters_extracted"] == len(_PARSED_TEXT)
    assert payload["size_bytes"] > 0
    assert len(payload["sha256"]) == 64
    # Job description fields
    assert payload["jd_filename"] == "job.txt"
    assert payload["jd_content_type"] == _TXT_CONTENT_TYPE
    assert payload["jd_extracted_text"] == _JOB_DESCRIPTION_TEXT
    assert payload["jd_characters_extracted"] == len(_JOB_DESCRIPTION_TEXT)
    assert payload["jd_size_bytes"] > 0
    assert len(payload["jd_sha256"]) == 64


async def test_upload_rejects_an_unsupported_cv_format(client: AsyncClient) -> None:
    response = await client.post(
        "/cv/upload",
        files={
            "cv_file": ("cv.txt", b"plain text", "text/plain"),
            "job_description_file": ("job.txt", _txt_bytes(), _TXT_CONTENT_TYPE),
        },
    )

    assert response.status_code == 400
    assert "docx" in response.json()["detail"]


async def test_upload_rejects_an_unsupported_job_description_format(client: AsyncClient) -> None:
    response = await client.post(
        "/cv/upload",
        files={
            "cv_file": ("cv.docx", _docx_bytes(), _DOCX_CONTENT_TYPE),
            "job_description_file": ("job.pdf", b"pdf content", "application/pdf"),
        },
    )

    assert response.status_code == 400
    assert "txt" in response.json()["detail"]


async def test_upload_rejects_a_cv_file_over_the_limit(client: AsyncClient) -> None:
    oversized = b"0" * (MAX_UPLOAD_SIZE_BYTES + 1)

    response = await client.post(
        "/cv/upload",
        files={
            "cv_file": ("cv.pdf", oversized, "application/pdf"),
            "job_description_file": ("job.txt", _txt_bytes(), _TXT_CONTENT_TYPE),
        },
    )

    assert response.status_code == 400
    assert "10 MB" in response.json()["detail"]


async def test_upload_requires_the_job_description_file(client: AsyncClient) -> None:
    response = await client.post(
        "/cv/upload",
        files={"cv_file": ("cv.docx", _docx_bytes(), _DOCX_CONTENT_TYPE)},
    )

    assert response.status_code == 422


async def test_upload_requires_a_cv_file(client: AsyncClient) -> None:
    response = await client.post(
        "/cv/upload",
        files={"job_description_file": ("job.txt", _txt_bytes(), _TXT_CONTENT_TYPE)},
    )

    assert response.status_code == 422


async def test_upload_end_to_end_with_the_real_parser() -> None:
    app.dependency_overrides.clear()
    client = _build_client()

    response = await client.post(
        "/cv/upload",
        files={
            "cv_file": ("cv.docx", _docx_bytes(), _DOCX_CONTENT_TYPE),
            "job_description_file": ("job.txt", _txt_bytes(), _TXT_CONTENT_TYPE),
        },
    )

    assert response.status_code == 200
    assert "Ana Developer" in response.json()["extracted_text"]
    assert _JOB_DESCRIPTION_TEXT in response.json()["jd_extracted_text"]


async def test_upload_end_to_end_with_a_real_pdf(sample_cv_pdf: bytes) -> None:
    app.dependency_overrides.clear()
    client = _build_client()

    response = await client.post(
        "/cv/upload",
        files={
            "cv_file": ("cv.pdf", sample_cv_pdf, "application/pdf"),
            "job_description_file": ("job.txt", _txt_bytes(), _TXT_CONTENT_TYPE),
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["filename"] == "cv.pdf"
    assert "Jhon Doe" in payload["extracted_text"]
    assert payload["characters_extracted"] > 1000
    assert _JOB_DESCRIPTION_TEXT in payload["jd_extracted_text"]


async def test_upload_rejects_a_pdf_without_a_text_layer() -> None:
    app.dependency_overrides.clear()
    client = _build_client()

    response = await client.post(
        "/cv/upload",
        files={
            "cv_file": ("scanned.pdf", _blank_pdf_bytes(), "application/pdf"),
            "job_description_file": ("job.txt", _txt_bytes(), _TXT_CONTENT_TYPE),
        },
    )

    assert response.status_code == 400
    assert "text" in response.json()["detail"].lower()