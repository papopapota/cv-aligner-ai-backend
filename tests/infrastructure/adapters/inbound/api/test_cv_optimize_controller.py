import io
from collections.abc import Generator

import docx
import pytest
from httpx import ASGITransport, AsyncClient

from src.application.services.optimize_cv_service import OptimizeCVService
from src.composition import build_optimize_cv_use_case
from src.domain.errors import AuditFailedError
from src.infrastructure.adapters.out.langgraph_orchestrator_adapter import (
    LangGraphOrchestratorAdapter,
)
from src.infrastructure.adapters.out.stub_llm_agent_adapter import (
    StubLLMAgentAdapter,
)
from src.main import app

_JOB_DESCRIPTION = "Buscamos Python developer con FastAPI y PostgreSQL."
_TXT_CONTENT_TYPE = "text/plain"
_DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _docx_bytes() -> bytes:
    document = docx.Document()
    document.add_paragraph("Ana Martinez")
    document.add_paragraph("Python developer con 6 anos de experiencia.")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _build_client() -> AsyncClient:
    return AsyncClient(transport=ASGITransport(app=app), base_url="http://test")


@pytest.fixture(autouse=True)
def _stub_the_llm() -> Generator[None, None, None]:
    app.dependency_overrides[build_optimize_cv_use_case] = lambda: (
        OptimizeCVService(
            LangGraphOrchestratorAdapter(
                StubLLMAgentAdapter(
                    variance_score=0.1,
                    variance_score_on_first_attempt=0.9,
                )
            )
        )
    )
    yield
    app.dependency_overrides.clear()


@pytest.fixture
def client() -> AsyncClient:
    return _build_client()


async def test_optimize_returns_the_audited_cv(client: AsyncClient) -> None:
    response = await client.post(
        "/cv/optimize",
        files={
            "cv_file": ("cv.docx", _docx_bytes(), _DOCX_CONTENT_TYPE),
            "job_description_file": ("job.txt", _JOB_DESCRIPTION.encode(), _TXT_CONTENT_TYPE),
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["optimized_content"]
    assert payload["variance_score"] == 0.1
    assert payload["writer_attempts"] == 2
    assert payload["is_within_threshold"] is True


async def test_optimize_with_the_real_pdf_fixture(
    client: AsyncClient,
    sample_cv_pdf: bytes,
) -> None:
    response = await client.post(
        "/cv/optimize",
        files={
            "cv_file": ("cv.pdf", sample_cv_pdf, "application/pdf"),
            "job_description_file": ("job.txt", _JOB_DESCRIPTION.encode(), _TXT_CONTENT_TYPE),
        },
    )

    assert response.status_code == 200
    assert response.json()["is_within_threshold"] is True


async def test_optimize_rejects_an_unsupported_cv_format(client: AsyncClient) -> None:
    response = await client.post(
        "/cv/optimize",
        files={
            "cv_file": ("cv.txt", b"plain", _TXT_CONTENT_TYPE),
            "job_description_file": ("job.txt", _JOB_DESCRIPTION.encode(), _TXT_CONTENT_TYPE),
        },
    )

    assert response.status_code == 400


async def test_optimize_rejects_a_non_txt_job_spec(client: AsyncClient) -> None:
    response = await client.post(
        "/cv/optimize",
        files={
            "cv_file": ("cv.docx", _docx_bytes(), _DOCX_CONTENT_TYPE),
            "job_description_file": ("job.pdf", b"%PDF-1.7 nope", "application/pdf"),
        },
    )

    assert response.status_code == 400
    assert "txt" in response.json()["detail"].lower()


async def test_optimize_requires_both_files(client: AsyncClient) -> None:
    response = await client.post(
        "/cv/optimize",
        files={"cv_file": ("cv.docx", _docx_bytes(), _DOCX_CONTENT_TYPE)},
    )

    assert response.status_code == 422


async def test_optimize_returns_422_when_the_variance_never_converges(
    client: AsyncClient,
) -> None:
    app.dependency_overrides[build_optimize_cv_use_case] = lambda: (
        OptimizeCVService(
            LangGraphOrchestratorAdapter(StubLLMAgentAdapter(variance_score=0.9))
        )
    )

    response = await client.post(
        "/cv/optimize",
        files={
            "cv_file": ("cv.docx", _docx_bytes(), _DOCX_CONTENT_TYPE),
            "job_description_file": ("job.txt", _JOB_DESCRIPTION.encode(), _TXT_CONTENT_TYPE),
        },
    )

    assert response.status_code == 422
    assert "variance audit" in response.json()["detail"]


async def test_the_audit_failure_is_a_business_error_not_a_server_error() -> None:
    assert issubclass(AuditFailedError, Exception)
    assert not issubclass(AuditFailedError, ConnectionError)
