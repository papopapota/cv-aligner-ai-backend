import asyncio
import io
from pathlib import Path

import docx
from docx.document import Document as DocxDocument
from pypdf import PdfReader

from src.application.ports.out.cv_parser_port import CVParserPort
from src.domain.errors import InvalidUploadError, UnsupportedFormatError

_PDF_EXTENSION = ".pdf"
_DOCX_EXTENSION = ".docx"


class DocumentParserAdapter(CVParserPort):
    async def parse(self, filename: str, content: bytes) -> str:
        return await asyncio.to_thread(self._parse_sync, filename, content)

    def _parse_sync(self, filename: str, content: bytes) -> str:
        extension = Path(filename).suffix.lower()
        if extension == _PDF_EXTENSION:
            return self._parse_pdf(content)
        if extension == _DOCX_EXTENSION:
            return self._parse_docx(content)
        raise UnsupportedFormatError(
            f"Unsupported file format '{extension}'. Supported: .pdf, .docx."
        )

    def _parse_pdf(self, content: bytes) -> str:
        try:
            reader = PdfReader(io.BytesIO(content))
            pages = [page.extract_text() or "" for page in reader.pages]
        except Exception as exc:
            raise InvalidUploadError(f"Could not read the PDF file: {exc}") from exc
        return "\n".join(page.strip() for page in pages if page.strip())

    def _parse_docx(self, content: bytes) -> str:
        try:
            document = docx.Document(io.BytesIO(content))
        except Exception as exc:
            raise InvalidUploadError(f"Could not read the DOCX file: {exc}") from exc
        return self._extract_docx_text(document)

    def _extract_docx_text(self, document: DocxDocument) -> str:
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(part for part in parts if part.strip())